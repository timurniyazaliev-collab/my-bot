import datetime

import openpyxl

HEADER_MAP = {
    "Сумма заказов по розничным ценам с учётом согласованной скидки, руб.": "order_sum",
    "Заказано, шт.": "order_qty",
    "Сумма продаж по розничным ценам с учётом согласованной скидки, руб.": "sale_sum",
    "Процент выкупа": "buyout_pct",
    "К перечислению за товар, руб.": "for_pay",
    "Стоимость логистики, руб.": "logistics",
    "Стоимость хранения, руб.": "storage",
}


def parse_summary_report(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

    col_idx = {}
    for name, key in HEADER_MAP.items():
        if name in headers:
            col_idx[key] = headers.index(name)

    by_date = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        day_cell = row[2]  # column C = День
        if not isinstance(day_cell, datetime.datetime):
            continue
        d = day_cell.date()
        by_date[d.isoformat()] = {key: row[idx] or 0 for key, idx in col_idx.items()}
    return by_date


def format_preview(by_date: dict) -> str:
    if not by_date:
        return "В файле не нашлось ни одной строки с датой по дням — проверь, что в отчёте включена разбивка по дням."

    lines = ["Черновик из твоего файла (пока никуда не записано):", ""]
    for key in sorted(by_date):
        v = by_date[key]
        d = datetime.date.fromisoformat(key)
        lines.append(
            f"{d.strftime('%d.%m')}: заказы {round(v.get('order_sum', 0))}₽ ({v.get('order_qty', 0)} шт) · "
            f"продажи {round(v.get('sale_sum', 0))}₽ · выкуп {round(v.get('buyout_pct', 0))}% · "
            f"к перечислению {round(v.get('for_pay', 0))}₽ · логистика {round(v.get('logistics', 0))}₽ · "
            f"хранение {round(v.get('storage', 0))}₽"
        )
    lines.append("")
    lines.append("Если цифры совпадают с тем, что ты видишь в кабинете WB — напиши /wb_apply, чтобы записать в таблицу.")
    return "\n".join(lines)
